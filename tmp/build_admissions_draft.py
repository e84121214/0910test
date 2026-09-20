from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\CCAM144F\Desktop\專題\0910test")
OUT = ROOT / "書審資料_母版初稿.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    paragraph.add_run(" 頁")


def apply_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = "Noto Sans CJK TC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Noto Sans CJK TC")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Noto Sans CJK TC")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, first_line=True, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        apply_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        apply_run_font(r2)
    else:
        r = p.add_run(text)
        apply_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        for run in p.runs:
            apply_run_font(run)
        if not p.runs:
            apply_run_font(p.add_run(item))
        else:
            p.runs[0].text = item


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        apply_run_font(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_cell_shading(c, "1F4E78")
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        apply_run_font(r, size=9.5, bold=True, color="FFFFFF")
        if widths:
            c.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(c, "F3F7FA")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if len(str(value)) < 20 and i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            apply_run_font(r, size=9)
            if widths:
                c.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Noto Sans CJK TC"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)

for name, size, bold, after in [
    ("Title", 24, True, 18),
    ("Subtitle", 12, False, 10),
    ("Heading 1", 16, True, 8),
    ("Heading 2", 13, True, 6),
    ("Heading 3", 11, True, 4),
]:
    style = styles[name]
    style.font.name = "Noto Sans CJK TC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(after)

footer = sec.footer.paragraphs[0]
apply_run_font(footer.add_run("研究所推甄書審資料母版初稿    "), size=8, color="666666")
add_page_number(footer)
for run in footer.runs:
    apply_run_font(run, size=8, color="666666")

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
r = p.add_run("研究所推甄書審資料母版初稿")
apply_run_font(r, size=24, bold=True)

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("環境時空資料與智慧決策")
apply_run_font(r, size=13)

doc.add_paragraph()
cover_rows = [
    ("申請者", "姓名待確認"),
    ("就讀學校", "國立成功大學"),
    ("就讀科系", "水利及海洋工程學系"),
    ("主要申請方向", "土木電腦輔助工程  水利與環境資訊  跨域智慧運算"),
    ("文件用途", "依各校簡章與研究方向客製前的內容母版"),
]
table = doc.add_table(rows=0, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_borders(table, color="E6E6E6")
for label, value in cover_rows:
    cells = table.add_row().cells
    cells[0].width = Cm(4)
    cells[1].width = Cm(11)
    set_cell_shading(cells[0], "F2F2F2")
    for c in cells:
        set_cell_margins(c, top=130, bottom=130)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    r = cells[0].paragraphs[0].add_run(label)
    apply_run_font(r, bold=True)
    r = cells[1].paragraphs[0].add_run(value)
    apply_run_font(r)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
r = p.add_run("本稿整合目前已確認的學業 研究 實習與課外經驗\n作為後續依系所客製與補件的基礎")
apply_run_font(r, size=10, color="555555")

doc.add_page_break()

add_heading(doc, "文件定位與核心主張", 1)
add_body(doc, "我希望在研究所階段持續處理環境資料中的時間、空間與不確定性問題，並學習如何把分析結果轉化為可供工程判斷使用的資訊。目前最能代表我的三項經驗，分別是台大土木電腦輔助工程組暑期實習中的氣象站圖建構與雷達類比驗證、以多來源資料預測霧事件的 GRU 實驗，以及成大實驗室研究助理工作中的淹水影像結構化判讀。")
add_body(doc, "這些專題的結果並非全部理想。雷達 analogue 的整體平均尚未穩定超過 useful FSS；GRU 在未知測站的泛化能力仍弱，且 persistence 在目前設定下仍具有更高 CSI；淹水影像 VLM 也仍屬概念驗證。這些限制讓我更確定，研究價值不在於套用更複雜的模型，而在於建立合理的資料切分、基準方法、評估指標與錯誤分析，並清楚界定模型可以支持的結論。")
add_body(doc, "因此，本份母版將我的核心形象定位為具有水利與海洋工程背景、能運用資料與人工智慧方法處理環境問題，並重視可信評估與實務使用邊界的申請者。後續投遞時，將依不同系所調整研究重點，而不改變已確認的事實與研究限制。")

add_heading(doc, "教授快速閱讀摘要", 2)
add_table(
    doc,
    ["審查問題", "目前可提供的回答"],
    [
        ("我是誰", "成大水利及海洋工程學生，以環境時空資料與智慧決策為跨域方向"),
        ("做過什麼", "完成台大 CAE 暑期實習專題，建置 GRU 霧事件比較流程，並參與 RAG 與 VLM 研究助理工作"),
        ("具備什麼能力", "真實資料清理、時序切分、空間驗證、基準比較、稀少事件評估、版本與環境管理"),
        ("研究態度", "保留負面結果，檢查資料與評估條件，不把局部高分包裝成整體成功"),
        ("為何讀研", "希望補強數學與計算基礎，進一步研究多來源環境資料融合與時空泛化"),
        ("仍需補強", "正式資訊核心課程與程式碼細節掌握度，並將階段成果整理成可公開的研究產出"),
    ],
    widths=[4.2, 12.6],
)

doc.add_page_break()

add_heading(doc, "個人履歷摘要", 1)
add_heading(doc, "學歷與學業概況", 2)
add_table(
    doc,
    ["項目", "內容"],
    [
        ("學校科系", "國立成功大學 水利及海洋工程學系"),
        ("修課進度", "截至 114 學年度第 2 學期累計 122.5 學分"),
        ("近期表現", "最近一學期 GPA 4.02，學期排名前 20%"),
        ("相關基礎", "工程數學 工程應用數學 遙測與 GIS 水資源及水利專業課程"),
        ("學業說明", "整體表現後期回升；正式累積平均與累積排名仍以校方證明為準"),
    ],
    widths=[4.2, 12.6],
)

add_heading(doc, "研究與實務經驗", 2)
add_table(
    doc,
    ["期間與身分", "題目與內容", "可驗證成果"],
    [
        ("2026 年 7 月至 8 月\n台大土木 CAE 暑期實習", "基於圖神經網路與氣象站網絡之大氣型態分類；建構測站圖、比較建邊權重、搜尋歷史類比事件，並以雷達回波與 FSS 驗證空間相似度", "實習證書 學術海報 四頁簡報 階段成果圖"),
        ("專題研究\nGRU 能見度與霧事件預測", "比較衛星 地面氣象與混合輸入；建立共同樣本、保留月份測試、C48 留站測試、F1 門檻與 persistence 參考基準", "三組可重現程式 結果表 metadata 訓練紀錄 門檻曲線"),
        ("成大實驗室研究助理", "參與 RAG 與 VLM 相關工作，將淹水照片判讀轉為可供流程串接的結構化欄位", "工作流程原型；正式計畫名稱 期間與個人分工待確認"),
    ],
    widths=[3.7, 8.7, 4.6],
)

add_heading(doc, "工具與方法", 2)
add_bullets(doc, [
    "資料與模型  Python NumPy pandas h5py GeoPandas PyTorch HDF5 CSV 與多測站時序資料",
    "研究方法  滑動視窗 連續區段 時間切分 標準化 缺值處理 Early Stopping 類別不平衡與門檻選擇",
    "評估方法  FSS useful FSS Precision Recall F1 CSI weighted BCE 與 persistence baseline",
    "工程流程  Git GitHub uv Conda Windows 與 Linux 遠端環境 長時間工作執行與結果留存",
])

add_heading(doc, "課外經驗", 2)
add_bullets(doc, [
    "成大流行音樂社公關  參與外部聯繫 演出協調 宣傳排程 器材需求整合 活動檢討與主唱課教學",
    "元大銀行第三屆數位金融校園大使  接觸金融科技 品牌任務與校園溝通；個人任務與成果待補",
    "成大鳳凰樹文學獎散文佳作  可作為文字組織與資訊轉譯能力的輔助證據",
    "第 21 屆高中生人文及社會科學營與高中學生媒體經驗  作為跨域好奇心與早期公關行政經驗的背景素材",
])

doc.add_page_break()

add_heading(doc, "自傳初稿", 1)
add_heading(doc, "從水利問題走向環境資料分析", 2)
add_body(doc, "就讀水利及海洋工程學系後，我逐漸發現，許多環境問題的困難不只來自方程式或模型，也來自資料本身的缺漏、尺度差異與空間分布。當衛星、氣象站、雷達與影像提供不同觀測時，真正困難的是如何讓它們在合理的時間與空間條件下被比較，並讓結果能支持工程判斷。這個問題成為我投入環境資料分析與人工智慧研究的起點。")

add_heading(doc, "在台大 CAE 實習學會檢驗方法的適用範圍", 2)
add_body(doc, "2026 年暑假，我參與台大土木電腦輔助工程組的暑期實習，專題為基於圖神經網路與氣象站網絡之大氣型態分類。研究先將氣象站視為節點，使用氣溫、氣壓與風場建立圖資料，比較 Distance、Correlation 與 Hybrid 三種建邊權重，再透過遮罩重建檢查模型能否表達測站間的空間關係。接著，我們利用高維特徵搜尋歷史類比事件，並以 QPESUMS 雷達回波和 Fractions Skill Score 檢驗降雨空間型態是否相似。")
add_body(doc, "結果讓我看到兩個需要謹慎處理的問題。第一，Distance 權重在目前的氣壓、氣溫與風場重建上較穩定，但 Hybrid 並未穩定優於純距離建邊；第二，部分最佳 analogue 可以超過 useful FSS，整體平均卻尚未支持方法已普遍有效。這表示候選事件中可能存在有參考價值的案例，但現有排序仍無法穩定分離好壞。這段經驗使我學會區分局部案例與整體結論，也理解評估條件不同時，不能直接比較絕對數值。")

add_heading(doc, "用 GRU 實驗理解基準模型與空間泛化", 2)
add_body(doc, "在能見度研究中，我將問題由單站回歸逐步改寫為多測站的霧事件分類，並比較純衛星、純地面氣象與混合輸入。為了讓比較具有意義，三個模型使用相同序列、標籤與切分，標準化參數只由訓練資料估計，分類門檻也只在 validation 上選擇。我另外保留月份作為 temporal test，並將 C48 完全排除於訓練之外，作為 spatial test。")
add_body(doc, "在保留月份測試中，純地面模型取得三個 GRU 中最高的 CSI 32.20%，混合模型則有最高 Recall 64.76%；然而 persistence 的 CSI 仍為 36.84%。到了 C48 留站測試，三個 GRU 的 CSI 僅介於 2.94% 至 5.91%，顯示未知測站泛化仍是主要瓶頸。這些結果沒有讓我停止研究，反而使問題變得更明確。我開始檢查霧事件在測站間高度不均衡、滑動視窗重疊、門檻校準與空間表示不足等限制，也理解更複雜的模型不能取代公平的基準與嚴格的測試設計。")

add_heading(doc, "將影像判讀轉為可使用的資訊", 2)
add_body(doc, "目前我在成大實驗室擔任由研究計畫經費支薪的研究助理，工作與 Retrieval Augmented Generation 及 Vision Language Model 有關。淹水影像原型的任務，是讓民眾上傳照片後，將是否可用、是否可見淹水、估計水深、流速等級、受困情況與危險因子轉為結構化欄位。我曾處理模型輸出的布林值型態與工作流程 schema 不一致等問題，也思考如何讓輸出串接資料庫、警示或人工覆核。")
add_body(doc, "這項原型仍不能取代現地感測與專業判斷。照片中的水深會受到透視、遮蔽與鏡頭角度影響，模型結果應定位為災情分流的輔助資訊。這段經驗讓我把研究關注從模型是否能回答，延伸到輸出能否被系統安全使用。")

add_heading(doc, "研究以外的協作與轉譯", 2)
add_body(doc, "研究之外，我曾參與流行音樂社的對外聯繫、演出協調、宣傳排程、器材需求整理與主唱課教學。與場地方或不同樂團合作時，我需要把零散需求整理成曲目、舞台配置、設備、時程與交付格式；活動結束後，也會記錄控時、人力、雨備與器材安排的問題。這些經驗訓練我把模糊需求轉成可執行規格，也使我在研究中習慣保留資料設定、版本與錯誤脈絡。")
add_body(doc, "我目前仍有需要補強之處。我的正式資訊核心修課較少，部分研究程式由 Codex 協助撰寫，我對研究問題、資料輸出入、清理、切分與實驗流程的掌握，仍強於逐行解釋所有實作細節。這是我在申請與口試前必須正面處理的能力缺口。我希望透過重寫最小版本、導讀關鍵函式與練習說明張量形狀、loss 和評估流程，使自己能更完整地為研究設計與程式實作負責。")
add_body(doc, "研究所階段，我希望進一步學習時空資料建模、多來源環境資料融合與可信評估，並把水利與氣象領域知識轉化為可被驗證、解釋與實際使用的決策資訊。")

doc.add_page_break()

add_heading(doc, "研究經驗一 氣象站圖建構與雷達類比驗證", 1)
add_heading(doc, "研究背景與任務", 2)
add_body(doc, "本研究源自台大土木電腦輔助工程組暑期實習。目標是使用氣象站網絡描述大氣狀態，從歷史資料中找出與目標事件相近的 analogue，再以雷達回波檢查候選事件的降雨空間型態是否具有參考價值。")

add_heading(doc, "方法與個人可說明內容", 2)
add_numbered(doc, [
    "將氣象站視為圖節點，使用氣溫、氣壓與風場資料建立輸入，並比較 Distance、Correlation 與 Hybrid 三種邊權重。",
    "透過遮罩部分節點或時段，評估模型重建測站氣象變數的能力，再把大氣狀態轉為高維特徵向量。",
    "為每個 target 搜尋 20 個歷史 analogue，使用 QPESUMS MREF3D21L 雷達產品比較 target 前後約一小時的降雨空間型態。",
    "以 30 dBZ 門檻與不同鄰域尺度計算 FSS，同時區分單一配對與事件平均，避免不同統計層級混用。",
])

add_heading(doc, "目前結果與研究限制", 2)
add_table(
    doc,
    ["觀察", "目前可支持的解讀"],
    [
        ("約 3841 筆結果，多數 FSS 落在 0 至 0.4", "整體相似度仍有限，不能宣稱 analogue 方法已普遍有效"),
        ("最大 FSS 約 0.9669", "個別候選事件可能具有高度空間相似性，但不能用最大值代表整體"),
        ("約 2080 筆為 NaN 或無法計算", "必須檢查事件稀少、有效網格、空間對齊與無事件等可能原因，尚不能指定單一成因"),
        ("部分 top analogue 超過 useful FSS，summary 未超過", "值得研究的問題是如何改善候選排序與篩選，而非只保留最高分案例"),
        ("Distance 權重較穩定，Hybrid 未穩定更好", "邊權重設計需要以任務與評估條件檢查，複合方法不一定自然優於簡單方法"),
    ],
    widths=[6.2, 10.6],
)

add_heading(doc, "此經驗形成的能力", 2)
add_bullets(doc, [
    "把氣象問題轉成可計算的圖結構與空間驗證問題",
    "處理 HDF5 時序檔案 雷達網格 有效值交集與多層級統計",
    "閱讀結果時區分最大值 平均值 top k 與缺值分布",
    "在結果未達預期時保留限制並提出可驗證的下一步問題",
])

doc.add_page_break()

add_heading(doc, "研究經驗二 多來源 GRU 霧事件預測", 1)
add_heading(doc, "研究設計", 2)
add_body(doc, "本實驗以過去 18 筆、每 10 分鐘一筆的資料，預測最後輸入後 60 分鐘是否出現能見度低於 1 公里的霧。三個模型採相同單層 GRU、64 hidden units、資料切分與隨機種子，只改變輸入來源，讓比較聚焦於衛星、地面氣象與混合資訊的差異。")
add_table(
    doc,
    ["資料切分", "序列數", "霧標籤數", "用途"],
    [
        ("Train", "140290", "1992", "25 站，用於訓練與估計標準化參數"),
        ("Validation", "37515", "489", "25 站，用於 early stopping 與選擇 F1 門檻"),
        ("Temporal test", "38744", "437", "已見測站的保留月份測試"),
        ("Spatial test C48", "5782", "94", "完全未參與訓練測站的空間泛化測試"),
    ],
    widths=[4.0, 3.0, 3.0, 6.8],
)

add_heading(doc, "主要結果", 2)
add_table(
    doc,
    ["模型", "Validation 門檻", "Temporal Precision", "Temporal Recall", "Temporal CSI", "C48 CSI"],
    [
        ("純衛星", "0.50", "32.25%", "45.54%", "23.27%", "2.94%"),
        ("純地面", "0.70", "41.69%", "58.58%", "32.20%", "3.97%"),
        ("衛星加地面", "0.61", "33.10%", "64.76%", "28.05%", "5.91%"),
        ("Persistence", "不適用", "53.24%", "54.46%", "36.84%", "38.64%"),
    ],
    widths=[3.6, 2.6, 2.8, 2.8, 2.8, 2.2],
)
add_body(doc, "純地面模型在三個 GRU 中有最高的 temporal precision 與 CSI，混合模型有最高 recall。Persistence 使用當下實測能見度，與三個 GRU 的輸入條件不同，因此不是完全公平的特徵消融比較；但它仍指出短期能見度延續具有很強訊號。C48 結果則顯示，目前模型無法可靠外推到未知測站。")

add_heading(doc, "研究反思", 2)
add_bullets(doc, [
    "霧樣本只占約 1% 至 1.6%，Accuracy 會被大量非霧樣本主導，因此以 Precision Recall CSI 與 weighted BCE 判讀。",
    "Train 的霧樣本集中於少數測站，模型可能學到站點偏差；單一 C48 也不足以估計所有新站的平均表現。",
    "滑動視窗高度重疊，序列數不能直接視為獨立事件數；後續不確定性應以事件或日期群組處理。",
    "門檻調整不足以修復 C48 的排序能力，下一步應優先檢查多站留出驗證 空間表示 站點不平衡與機率校準。",
])

doc.add_page_break()

add_heading(doc, "研究經驗三 淹水影像結構化判讀", 1)
add_heading(doc, "研究情境", 2)
add_body(doc, "在成大實驗室研究助理工作中，我參與 RAG 與 VLM 相關任務。淹水影像原型希望將民眾上傳的照片轉為可以進入後續系統的結構化資訊，降低人工初步分流時需要反覆閱讀與整理的成本。")

add_heading(doc, "目前原型欄位", 2)
add_table(
    doc,
    ["類別", "欄位與功能"],
    [
        ("影像品質", "image usable  判斷照片是否足以進行後續分析"),
        ("淹水狀態", "flood visible  判斷是否可見淹水"),
        ("水深估計", "depth min cm  depth max cm  depth basis  保留估計區間與依據"),
        ("流況", "flow level  提供粗略流速或流況分級"),
        ("受困與危險", "people trapped visible  vehicle trapped visible  underground or underpass  electrical hazard"),
    ],
    widths=[4.2, 12.6],
)

add_heading(doc, "工程與風險意識", 2)
add_body(doc, "工作流程曾遇到模型以字串 true 與 false 表示布林值的問題，需要在 code node 中轉換型態，並確認節點 schema 與介面顯示是否一致。這個細節使我理解，模型給出答案只是流程的一部分；若資料型態與欄位定義不穩定，後續資料庫或警示服務仍可能失效。")
add_body(doc, "原型目前只能作為輔助分流。影像水深估計會受到透視、遮蔽、參考物與拍攝角度影響，VLM 也不應取代現場感測與專業人員。後續研究需要建立資料品質分級、不確定性表達、人工覆核與錯誤案例分析。")

add_heading(doc, "三項研究的共同脈絡", 2)
add_table(
    doc,
    ["共同問題", "在三項研究中的呈現"],
    [
        ("資料尺度與品質", "雷達網格有效值  多站時序缺漏  影像可用性"),
        ("時間與空間結構", "analogue 時間窗與 FSS 鄰域  GRU 保留月份與留站測試  影像中的場景與位置條件"),
        ("可信評估", "useful FSS  persistence baseline  structured output 與人工覆核"),
        ("實務界線", "不以 top case 代表整體  不宣稱未知站泛化  不把原型稱為正式防災系統"),
    ],
    widths=[4.2, 12.6],
)

doc.add_page_break()

add_heading(doc, "讀書與研究計畫初稿", 1)
add_heading(doc, "入學前補強", 2)
add_numbered(doc, [
    "完成研究程式導讀與最小版本重寫，能口頭說明資料流、張量形狀、loss、訓練迴圈、門檻選擇與評估流程。",
    "補強線性代數、機率統計、資料結構與演算法等計算基礎，並以目前專題中的實際問題驗證理解。",
    "整理台大 CAE 實習、GRU 與 VLM 的版本、圖表、程式與個人分工，將可公開範圍與研究限制寫成一致的研究紀錄。",
    "針對目標系所閱讀近期教師論文與課程規劃，確認既有經驗與研究所階段欲補強能力的具體連結。",
])

add_heading(doc, "研究所前期", 2)
add_body(doc, "研究所前期將以數學、統計、機器學習與空間資料分析為核心，補足跨領域訓練中較不完整的理論基礎。同時，我希望把既有環境資料專題整理成可重現的基準流程，包含資料版本、切分、特徵、模型、門檻、指標與錯誤案例，使後續方法改進有可靠比較基礎。實際課名與修課順序需依目標系所課程確認。")

add_heading(doc, "研究所中後期", 2)
add_body(doc, "研究方向暫定聚焦於多來源環境時空資料融合與未知站泛化。方法上可比較傳統統計與基準模型、GRU 等時序模型、圖神經網路或圖卷積循環模型，並檢查不同模型在跨月份、跨測站與稀少事件下的表現。若資料條件允許，也將評估物理先驗、站點拓撲、機率校準與不確定性估計，而不是只追求單一測試集上的最高分數。")
add_body(doc, "成果評估將至少包含三個層次。第一，模型是否在與簡單基準相同或明確說明差異的條件下改善；第二，結果是否能跨事件、月份或站點維持；第三，輸出是否具備可解釋的錯誤範圍，能被工程流程或決策者合理使用。")

add_heading(doc, "預期能力與發展", 2)
add_bullets(doc, [
    "建立能處理多來源環境資料的時空模型與評估流程",
    "能從領域問題決定資料 表示 模型與指標，而非只套用現成架構",
    "將研究結果轉化為可追蹤 可解釋並保留不確定性的決策資訊",
    "形成能獨立閱讀論文 重現方法 分析失效條件並清楚溝通的研究能力",
])

doc.add_page_break()

add_heading(doc, "依申請方向調整的動機草稿", 1)
add_heading(doc, "土木電腦輔助工程方向", 2)
add_body(doc, "我希望把土木與環境問題中的時間、空間與網絡結構，轉化為可被計算模型學習與檢驗的表示。台大 CAE 暑期實習讓我實際接觸氣象站建圖、圖神經網路與空間驗證，也使我看到自己在計算理論與程式實作細節上的不足。申請此方向時，我會以實習與多站 GRU 作為核心證據，說明我已能定義研究問題與評估設計，並希望進一步補強圖機器學習、科學運算與工程資訊整合能力。正式版本仍需依當年度課程、教授近期論文與實驗室方向客製。")

add_heading(doc, "水利與環境資訊方向", 2)
add_body(doc, "我希望延續水利及海洋工程訓練，深入研究雷達、衛星、地面測站與影像等資料如何共同描述環境事件。FSS 研究與霧事件預測讓我看到，空間尺度、季節切分、站點分布與基準方法會直接改變結論；淹水影像原型則讓我思考模型輸出如何進入防災流程。申請水利相關方向時，我會強調環境資料的物理意義、研究方法與決策使用邊界，並正面說明部分水力核心課程成績不突出，以及後期學業回升與研究實作如何支持進一步學習。")

add_heading(doc, "跨域智慧運算方向", 2)
add_body(doc, "我從水利與氣象問題進入人工智慧，不具備傳統資訊科系完整的核心修課，但累積了真實資料清理、模型比較、系統流程與研究限制分析的經驗。申請跨域智慧運算方向時，我會把領域資料與可信評估作為差異化基礎，並將線性代數、資料結構、演算法與程式碼掌握度列為明確補強目標。對資訊或電機等高風險志願，正式版本必須鎖定高度吻合的資料與 AI 研究方向，不能只以工具名稱取代基礎能力證明。")

doc.add_page_break()

add_heading(doc, "課外經驗與個人特質", 1)
add_heading(doc, "把零散資訊整理成可執行規格", 2)
add_body(doc, "在流行音樂社擔任公關期間，我參與外部合作、演出協調、宣傳排程與器材需求整理。面對場地方、樂團、幹部與宣傳端不同的資訊，我需要確認曲目、舞台配置、麥克風、音箱、鼓組、導線與截止時間，再將內容轉成各方可以執行的規格。這種工作方式延伸到研究中，就是先確認欄位、時間基準、資料版本與輸出格式，再進行分析。")

add_heading(doc, "用復盤改善研究與活動", 2)
add_body(doc, "社團活動後的檢討會記錄了人力、控時、雨備、舞台、器材與音控問題；研究實驗則保留資料切分、baseline、門檻與錯誤分布。兩種情境都使我習慣回到失敗條件，找出下一輪可以改變的因素。當 GRU 未超過 persistence、C48 泛化表現偏低時，我沒有只保留較好的曲線，而是重新檢查資料與評估設計。這是目前最能代表我的工作特質。")

add_heading(doc, "教學與文字轉譯", 2)
add_body(doc, "參與主唱課規劃時，我需要把呼吸、共鳴、真假音、混聲與歌曲詮釋拆成循序內容；公關文案則需要依社員、觀眾與合作方調整資訊密度。大一獲得鳳凰樹文學獎散文佳作，也提供了文字表達的外部證據。這些經驗使我在說明研究時，會先建立問題與直覺，再說明方法、數值與限制。")

add_heading(doc, "跨域經驗的使用原則", 2)
add_body(doc, "元大銀行校園大使、人文社會科學營與高中學生媒體經驗可作為跨域好奇心與溝通背景，但目前部分任務、分工與成果尚未確認。正式書審將只保留能支持申請主軸且有證據的內容，不平均羅列所有參與經歷。")

doc.add_page_break()

add_heading(doc, "審查風險與補強說明", 1)
add_table(
    doc,
    ["可能疑慮", "目前狀態", "書審與口試處理方式"],
    [
        ("資訊核心修課不足", "未見線性代數 離散數學 資料結構 演算法等完整正式修課", "不以工具名稱掩飾；用研究設計與實作證據說明現有能力，並提出具體補課與自學計畫"),
        ("水力核心成績不突出", "流體力學一 76 二 70 明渠水力學 67", "不為單科成績找藉口；呈現近期 GPA 4.02 與研究實作，並準備說明後續如何補足理論"),
        ("程式碼理解不完整", "掌握實驗架構 資料流 清洗切分與評估，部分程式由 Codex 協助", "不宣稱全數獨立手刻；完成最小版本重寫與關鍵函式導讀，準備逐段口頭說明"),
        ("模型結果未形成成功案例", "GRU 未超過 persistence，未知站泛化弱；FSS 整體平均亦有限", "以公平比較 錯誤分析與研究限制作為方法能力證據，不把負面結果改寫成成功"),
        ("研究身分與分工需更精確", "CAE 實習有證書；RA 正式計畫與分工仍待補", "補齊期間 計畫名稱 指導者 個人貢獻 共同工作與可公開範圍"),
    ],
    widths=[3.4, 6.0, 7.4],
)

add_heading(doc, "口試前應能回答的問題", 2)
add_numbered(doc, [
    "為什麼用 FSS，而不是只比較像素差或相關係數；useful FSS 的意義是什麼。",
    "為什麼 mask probability 不同的重建結果不能直接以絕對 R² 比較。",
    "GRU 的輸入與標籤如何對齊，為什麼序列不能跨越缺時、測站與月份。",
    "為什麼三個模型使用共同樣本，門檻只能由 validation 選擇。",
    "為什麼 persistence 並非完全公平的同輸入比較，卻仍是重要參考。",
    "C48 表現差可能來自哪些因素，哪些只是推測，下一步如何設計實驗確認。",
    "VLM 的欄位型態與 schema 為何重要，系統如何處理不確定與人工覆核。",
    "哪些程式由自己理解與修改，哪些由 AI 協助，如何驗證程式沒有造成資料洩漏或評估錯誤。",
])

doc.add_page_break()

add_heading(doc, "現有證據與附件規劃", 1)
add_table(
    doc,
    ["素材", "可證明內容", "建議放置方式"],
    [
        ("正式中文成績單", "122.5 學分 最近一學期 GPA 4.02 與學期排名", "依簡章上傳；正文只摘要與申請方向有關的趨勢與課程"),
        ("台大 CAE 實習證書", "2026 年 7 月 1 日至 8 月 26 日 實習單位 專題名稱與表現優良", "履歷與研究經驗頁摘要，證書置附件"),
        ("CAE 海報與四頁簡報", "測站建圖 邊權重 遮罩重建 analogue 與 FSS 流程", "選一張研究架構與一張代表結果，避免整份投影片堆疊"),
        ("GRU 程式與結果檔", "共同樣本 切分 門檻 baseline 與三模型指標", "正文放比較表；程式與完整結果以作品連結或口試備查"),
        ("VLM 工作流程", "結構化欄位 型態處理與系統串接概念", "確認公開範圍後放流程圖與一個錯誤處理例子"),
        ("熱音社紀錄", "外部合作 宣傳排程 器材整合 教學與活動復盤", "只選能直接支持協作 轉譯或反思能力的紀錄"),
        ("鳳凰樹文學獎證明", "散文佳作與文字表達", "作為輔助證據，不取代研究成果"),
    ],
    widths=[4.0, 7.0, 5.8],
)

add_heading(doc, "送件前待確認資料", 2)
add_bullets(doc, [
    "姓名 聯絡方式 正式累積平均 累積排名 語言成績 獎學金 證照與其他可驗證成果",
    "每一目標系所當年度簡章 頁數 上傳欄位 截止日 欲申請教授與近期研究",
    "台大 CAE 實習中本人完成 共同完成與 AI 協助的範圍，以及推薦信可公開程度",
    "GRU 專題正式題名 指導者 團隊分工與最新可公開版本",
    "成大研究助理的正式計畫名稱 任職期間 教師系所與職稱 個人分工 交付成果與公開範圍",
    "元大校園大使 熱音社 人文社科營 高中學生媒體與鳳凰樹文學獎的年份 職稱 任務與證明",
    "研究所畢業後規劃，以及各目標系所對應的研究問題與選擇理由",
])

add_heading(doc, "這份母版的客製方式", 2)
add_body(doc, "正式送件時不會把本稿原封不動投遞。每一系所應依簡章重新分配篇幅，簡歷優先保留學業、研究、實習與直接相關能力；自傳負責說明研究方向如何形成；申請動機連結既有經驗、能力缺口、系所資源與未來問題；讀書與研究計畫則提出可執行但仍可調整的學習路徑。遮住校名後，內容仍應能看出為該系所撰寫。")

# Keep headings with following content and normalize all fonts.
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") or p.style.name in ("Title", "Subtitle"):
        p.paragraph_format.keep_with_next = True
    for run in p.runs:
        if run.font.name is None:
            apply_run_font(run)

doc.core_properties.title = "研究所推甄書審資料母版初稿"
doc.core_properties.subject = "環境時空資料與智慧決策"
doc.core_properties.author = ""
doc.core_properties.keywords = "研究所推甄 書審 水利 氣象 人工智慧"
doc.save(OUT)
print(OUT)
